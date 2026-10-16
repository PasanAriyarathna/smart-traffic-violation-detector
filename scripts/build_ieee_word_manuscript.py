"""Build IEEE conference Word paper from template + repo evidence + local references only."""

from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(r"C:\Users\Pasan\OneDrive\Desktop\traffic\conference-template-letter.docx")
OUTPUT = ROOT / "docs/research_paper/Smart_Traffic_Violation_IEEE_Paper_Final.docx"
METRICS = ROOT / "docs/evaluation/detector/detector_val_test_20261009T120646Z/metrics.json"
PIPELINE_FIG = ROOT / "docs/research_paper/ieee/generated/figures/pipeline.png"

def _normalize_template_ooxml(src: Path, dest: Path) -> None:
    """Copy IEEE strict template and rewrite strict OOXML namespaces for python-docx."""
    ns_replacements = [
        (
            "http://purl.oclc.org/ooxml/officeDocument/relationships/",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/",
        ),
        (
            "http://purl.oclc.org/ooxml/wordprocessingml/main",
            "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        ),
        (
            "http://purl.oclc.org/ooxml/drawingml/main",
            "http://schemas.openxmlformats.org/drawingml/2006/main",
        ),
        (
            "http://purl.oclc.org/ooxml/drawingml/wordprocessingDrawing",
            "http://schemas.openxmlformats.org/drawingml/wordprocessingDrawing",
        ),
        (
            "http://purl.oclc.org/ooxml/officeDocument/math",
            "http://schemas.openxmlformats.org/officeDocument/2006/math",
        ),
    ]
    with zipfile.ZipFile(src, "r") as zin, zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.endswith((".rels", ".xml")):
                text = data.decode("utf-8")
                for old, new in ns_replacements:
                    text = text.replace(old, new)
                data = text.encode("utf-8")
            zout.writestr(item, data)


REFERENCES: list[str] = [
    'W. G. S. R. Thilakarathna et al., "Road traffic accidents in Sri Lanka: A retrospective analysis and artificial intelligence-based solutions for prevention," Sri Lanka J. Forensic Medicine, Science & Law, vol. 16, no. 1, pp. 22–29, Jun. 2025, doi: 10.4038/sljfmsl.v16i1.7991.',
    'O. R. Polu, "AI-powered traffic violation detection using CCTV footage," Int. J. Sci. Res., vol. 13, no. 3, Mar. 2024.',
    'C. S. Ong, T. Connie, and M. K. O. Goh, "Traffic violation detection using computer vision techniques," Int. J. Informatics Visualization, published Nov. 30, 2024.',
    'D. Dede et al., "Next-gen traffic surveillance: AI-assisted mobile traffic violation detection system," arXiv:2311.16179, Nov. 2023.',
    'S. R. Anand, N. Kilari, and D. U. S. Raj Kumar, "Traffic signal violation detection using artificial intelligence and deep learning," Int. J. Adv. Res. Eng. Technol., vol. 12, no. 2, pp. 207–217, Feb. 2021, doi: 10.34218/IJARET.12.2.2021.019.',
    'L. Q. Thao et al., "Automatic traffic red-light violation detection using AI," Ingénierie des Systèmes d\'Information, vol. 27, no. 1, doi: 10.18280/isi.270109, 2022.',
    'S. Asadianfam, M. Shamsi, and A. R. Kenari, "TVD-MRDL: Traffic violation detection system using MapReduce-based deep learning for large-scale data," Multimedia Tools Appl., doi: 10.1007/s11042-020-09714-8, 2020.',
    'S. K. Dwivedi, R. Amin, and S. Vollala, "Blockchain-based secured IPFS-enable event storage technique with authentication protocol in VANET," IEEE/CAA J. Autom. Sinica, vol. 8, no. 12, pp. 1914–1928, Dec. 2021.',
    'M. R. P. N. Bandara and X. Cheng, "An automatic number plate recognition system for Sri Lankan vehicle number plates," Int. J. Sci. Res. Eng. Develop., vol. 6, no. 3, pp. 76–82, May–Jun. 2023.',
    'A. R. M. Nizzad et al., "Internet of Things based automatic system for the traffic violation," in Proc. 5th Int. Conf. Elect., Electron., Commun., Comput. Technol. Optim. Techn. (ICEECCOT), Dec. 2021, pp. 371–376.',
    'S. T. Rakkesh, A. R. Weerasinghe, and R. A. C. Ranasinghe, "Simulation of real-time vehicle speed violation detection using complex event processing," Dept. Comput. Sci., Univ. Colombo Sch. Comput., Colombo, Sri Lanka, 2016.',
    'R. Bhakre, N. Madane, and M. Narkhede, "AI based traffic violation detection," Smt. Indira Gandhi Coll. Eng., India, 2021.',
    'A. Kumar et al., "S-TVDS: Smart traffic violation detection system for Indian traffic scenario," Int. J. Innovative Technol. Exploring Eng., vol. 8, no. 4S3, pp. 6–10, Mar. 2019, doi: 10.35940/ijitee.D1002.0384S319.',
]


def _clear_document_body(doc: Document) -> None:
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)


def _add(doc: Document, style: str, text: str) -> None:
    p = doc.add_paragraph(text, style=style)


def _add_cite(doc: Document, style: str, parts: list[str | int]) -> None:
    """parts: strings and int citation numbers."""
    p = doc.add_paragraph(style=style)
    for part in parts:
        if isinstance(part, int):
            p.add_run(f"[{part}]")
        else:
            p.add_run(part)


def _add_table_detection(doc: Document, m: dict, n: int) -> None:
    _add(doc, "tablehead", "TABLE I")
    _add(doc, "tablehead", "OBJECT DETECTION ON HELD-OUT TEST SPLIT (AUTHORITATIVE RUN)")
    table = doc.add_table(rows=5, cols=2)
    rows = [
        ("Test images", str(n)),
        ("Precision", f"{m['precision']:.3f}"),
        ("Recall", f"{m['recall']:.3f}"),
        ("mAP@0.5", f"{m['map50']:.3f}"),
        ("mAP@0.5:0.95", f"{m['map50_95']:.3f}"),
    ]
    for i, (a, b) in enumerate(rows):
        table.rows[i].cells[0].text = a
        table.rows[i].cells[1].text = b
    _add(
        doc,
        "BodyText",
        "Source: docs/evaluation/detector/detector_val_test_20261009T120646Z/metrics.json "
        "(CPU, conf=0.25, IoU=0.7, imgsz=640). These metrics measure annotated object classes only.",
    )


def _add_table_rules(doc: Document) -> None:
    _add(doc, "tablehead", "TABLE II")
    _add(doc, "tablehead", "RULE-LEVEL AND OCR METRIC AVAILABILITY (2026-10-09 REPOSITORY STATE)")
    table = doc.add_table(rows=5, cols=3)
    hdr = ["Metric", "Eligible labeled samples", "Status"]
    data = [
        ("OCR mean CER / exact-match", "0 / 30 plate crops", "NOT MEASURED"),
        ("Helmet rule P/R/F1", "0 TP/FP; 48 unreviewed", "NOT MEASURED"),
        ("Red-light event P/R/F1", "0 TP/FP; 1 unlabeled candidate", "NOT MEASURED"),
        ("Overtaking event P/R/F1", "0 predictions; no calibrated GT", "NOT MEASURED"),
    ]
    for j, h in enumerate(hdr):
        table.rows[0].cells[j].text = h
    for i, row in enumerate(data, start=1):
        for j, val in enumerate(row):
            table.rows[i].cells[j].text = val


def _add_table_related(doc: Document) -> None:
    _add(doc, "tablehead", "TABLE III")
    _add(doc, "tablehead", "COMPARISON OF PRIOR TRAFFIC-VIOLATION SYSTEMS (FROM LOCAL REFERENCE SET)")
    table = doc.add_table(rows=5, cols=4)
    table.rows[0].cells[0].text = "Study"
    table.rows[0].cells[1].text = "Detection / tracking"
    table.rows[0].cells[2].text = "Violation focus"
    table.rows[0].cells[3].text = "Reported limitation (from source)"
    rows = [
        (
            "Ong et al. [3]",
            "YOLO + Deep SORT",
            "Red-light, double-line overtaking",
            "CCTV-dependent geometry; no human-in-the-loop governance discussed",
        ),
        (
            "Dede et al. [4]",
            "YOLOv5 + strongSORT",
            "Six violation types (mobile dashcam)",
            "End-to-end violation notices; field labels not in our scope",
        ),
        (
            "Polu [2]",
            "YOLO + transformers (proposed)",
            "Multi-type CCTV violations",
            "Claims high accuracy on urban datasets; independent audit not reproduced here",
        ),
        (
            "This prototype",
            "YOLOv8n + ByteTrack (repo)",
            "Helmet, red-light, overtaking candidates",
            "Explicit NOT MEASURED rule metrics; NEEDS_REVIEW outputs",
        ),
    ]
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            table.rows[i].cells[j].text = val


def build_manuscript() -> None:
    if not TEMPLATE.is_file():
        raise SystemExit(f"Template missing: {TEMPLATE}")
    metrics_doc = json.loads(METRICS.read_text(encoding="utf-8"))
    m = metrics_doc["metrics"]
    n = metrics_doc["test_image_count"]

    with TemporaryDirectory() as tmp:
        normalized = Path(tmp) / "ieee_template_normalized.docx"
        _normalize_template_ooxml(TEMPLATE, normalized)
        doc = Document(str(normalized))
        _clear_document_body(doc)
        _write_manuscript_body(doc, metrics_doc, m, n)
        doc.save(str(OUTPUT))
    print(f"Wrote {OUTPUT}")


def _write_manuscript_body(doc: Document, metrics_doc: dict, m: dict, n: int) -> None:

    _add(
        doc,
        "papertitle",
        "Smart Traffic Violation Detection Using YOLOv8, Rule-Based Reasoning, and Human-in-the-Loop Review",
    )
    _add(
        doc,
        "Author",
        "H. A. P. P. Ariyarathna\n"
        "Faculty of Information Technology (Horizon Campus)\n"
        "Malabe, Sri Lanka\n"
        "ITBIN-2313-0009",
    )
    _add(
        doc,
        "Author",
        "A. S. N. M. Seeman\n"
        "Faculty of Information Technology (Horizon Campus)\n"
        "Malabe, Sri Lanka\n"
        "ITBIN-2313-0103",
    )

    abstract = (
        "We present a research prototype for Sri Lankan dashcam traffic analysis that combines a custom "
        "YOLOv8n detector (13 object classes), ByteTrack multi-object tracking, conservative geometric "
        "rule layers for helmet association, red-light stop-line crossing, and overtaking maneuvers, "
        "optional EasyOCR-based plate reading with SHA-256 hashing, human verification workflows, and a "
        "Solidity ViolationLogger contract tested on a local Hardhat network. All rule outputs are labeled "
        "NEEDS_REVIEW rather than confirmed legal violations. On a held-out image test split of "
        f"{n} frames (CPU, confidence 0.25, IoU 0.7), object detection achieved precision "
        f"{m['precision']:.3f}, recall {m['recall']:.3f}, mAP@0.5 {m['map50']:.3f}, and "
        f"mAP@0.5:0.95 {m['map50_95']:.3f}. These metrics measure annotated object-class detection only; "
        "they do not establish red-light or illegal-overtaking violation accuracy. Helmet rule precision, "
        "license-plate character error rate, and real-video red-light/overtaking precision, recall, and F1 "
        "are not measured pending human labels and verified signal/movement protocols. A preliminary "
        "single-clip investigation (track 130 at 11.788 s) illustrates candidate generation and documented "
        "limits of manual signal intervals versus visible signal aspects."
    )
    p = doc.add_paragraph(style="Abstract")
    p.add_run("Abstract—").italic = True
    p.add_run(abstract)

    p = doc.add_paragraph(style="Keywords")
    p.add_run("Keywords—").bold = True
    p.add_run(
        "Traffic violation detection, YOLOv8, human-in-the-loop review, dashcam video, "
        "geometric rules, blockchain metadata"
    )

    _add(doc, "Heading1", "Introduction")
    _add_cite(
        doc,
        "BodyText",
        [
            "Road traffic injuries remain a major public-health burden; retrospective forensic analysis in Sri Lanka "
            "links many fatalities to human factors and rule breaking, motivating supervised AI assistance rather than "
            "autonomous adjudication ",
            1,
            ". Prior systems combine deep learning detectors with tracking and heuristics for red-light jumping, "
            "speeding, lane discipline, and helmet compliance ",
            2,
            ", ",
            3,
            ", ",
            12,
            ". However, published pipelines often report end-to-end violation accuracy without separating object "
            "detection from legal violation classification or documenting human review ",
            3,
            ", ",
            5,
            ".",
        ],
    )
    _add_cite(
        doc,
        "BodyText",
        [
            "This IT41043 Intelligent Systems prototype detects traffic-related objects with a locally trained YOLOv8n "
            "model, tracks with ByteTrack inside model/violation_detector.py, applies opt-in rule modules that emit "
            "review candidates, exports JSONL events and evidence frames, supports manual review via FastAPI/SQLite, "
            "and optionally logs evidence hashes through an Ethereum-compatible logger. We report measured detection "
            "metrics, explicit NOT MEASURED gaps for rule-level evaluation, and a red-light case study that must not "
            "be interpreted as confirmed violation detection.",
        ],
    )

    _add(doc, "Heading1", "Problem Statement and Research Objectives")
    _add(
        doc,
        "BodyText",
        "Problem: Dashcam footage contains rich context, but converting detections into enforceable violation claims "
        "requires stop-line geometry, signal timing, movement-specific aspects, and human adjudication.",
    )
    _add(
        doc,
        "BodyText",
        "Objectives: (1) Train and evaluate a multi-class detector with reproducible metrics on the project test "
        "split. (2) Integrate tracking and rule layers that flag candidates with evidence artifacts. (3) Separate "
        "evaluation levels: object detection versus rule-event generation versus human-verified violation "
        "classification. (4) Document limitations including small test split, split-leakage risk, and unlabeled rule/OCR "
        "ground truth.",
    )
    _add(
        doc,
        "BodyText",
        "Research questions: RQ1—What object-detection quality is achieved on the held-out test split under fixed "
        "inference settings? RQ2—Can conservative geometric rules emit auditable candidates without claiming legal "
        "violations? RQ3—Which violation metrics remain unmeasurable until human labels and calibrated camera "
        "configuration exist?",
    )

    _add(doc, "Heading1", "Related Work and Research Gap")
    _add_cite(
        doc,
        "BodyText",
        [
            "Computer-vision pipelines for traffic enforcement commonly pair YOLO-family detectors with multi-object "
            "trackers. Ong et al. integrate YOLO with Deep SORT, automated lane/stop-line geometry, and point-line "
            "distance tests for red-light and double-line overtaking on CCTV footage ",
            3,
            ". Dede et al. describe a mobile dashcam system with YOLOv5 detection, strongSORT tracking, and six "
            "violation-specific algorithms ending in license-plate identification for notices ",
            4,
            ". Polu proposes CCTV analytics with YOLO detection, transformer-based action recognition, edge AI, "
            "federated learning, and blockchain-stored evidence ",
            2,
            ".",
        ],
    )
    _add_cite(
        doc,
        "BodyText",
        [
            "Red-light-focused studies emphasize ROI placement and signal-state timing: Thao et al. report high "
            "accuracy for vehicle, signal-state, and violation stages using modified YOLOv5s on COCO ",
            6,
            "; Anand et al. use YOLOv3 with ROI-based signal-jump logic ",
            5,
            ". Large-scale and IoT-oriented designs address scalability and identification—Asadianfam et al. use "
            "MapReduce deep learning on Hadoop ",
            7,
            "; Nizzad et al. combine IoT with image processing for speed violations in Sri Lanka ",
            10,
            ". Rakkesh et al. simulate speed-violation monitoring with complex event processing for Sri Lankan districts ",
            11,
            ", while Kumar et al. propose a smart traffic violation detection concept for dense urban Indian traffic ",
            13,
            ".",
        ],
    )
    _add_cite(
        doc,
        "BodyText",
        [
            "Integrity and plate reading intersect with enforcement workflows. Dwivedi et al. combine blockchain with "
            "IPFS for distributed event storage and vehicle authentication in VANET ",
            8,
            ". Bandara and Cheng design ANPR for Sri Lankan plates using OpenCV and Tesseract with reported plate "
            "detection accuracy on ten sample images ",
            9,
            ".",
        ],
    )
    _add(
        doc,
        "BodyText",
        "Gap addressed here: prior work often reports violation-detection accuracy without publishing independent "
        "human labels or separating detection mAP from event-level precision/recall. Our prototype deliberately "
        "exports NEEDS_REVIEW candidates, documents NOT MEASURED rule metrics, and keeps hash-based audit logging "
        "optional and environment-gated.",
    )
    _add_table_related(doc)

    _add(doc, "Heading1", "System Architecture and Methodology")
    _add(
        doc,
        "BodyText",
        "Pipeline: video input → YOLOv8 detection (Ultralytics) → ByteTrack (bytetrack.yaml) → optional rule modules "
        "→ annotated video + JSONL + evidence JPEGs → FastAPI import/review → optional chain client.",
    )
    if PIPELINE_FIG.is_file():
        doc.add_picture(str(PIPELINE_FIG), width=Inches(3.2))
        _add(doc, "figurecaption", "Fig. 1. End-to-end prototype workflow (conceptual; not performance evidence).")
    _add(
        doc,
        "BodyText",
        "Rule modules (opt-in via JSON): Helmet association links persons and motorcycles geometrically when no "
        "helmet box overlaps the rider head region. Red-light rule (model/rules/red_light.py) requires stop_line, "
        "approach_side, and manually merged signal_states intervals; crossing uses bbox center segment intersection; "
        "states are red, amber, green, or unknown only—no circular-versus-arrow aspect field and no pixel-level signal "
        "inference from traffic_light detections. Overtaking rule uses normalized lane/prohibition polygons; template "
        "configs with empty polygons block real-video evaluation until calibration.",
    )
    _add(
        doc,
        "BodyText",
        "OCR: optional EasyOCR on number_plate crops (model/plate_ocr.py). Blockchain: ViolationLogger.sol stores "
        "hashes; Hardhat tests verify access control. These components are implemented in the repository; Sepolia "
        "latency was not measured in the authoritative evaluation snapshot.",
    )

    _add(doc, "Heading1", "Experimental Setup")
    _add(
        doc,
        "BodyText",
        f"Object-detection evaluation uses weights documented in metrics.json (SHA-256 "
        f"{metrics_doc['config']['weights_sha256'][:16]}…) on data.yaml test split with "
        f"{n} images, device CPU, conf=0.25, IoU=0.7, imgsz=640. Rule-layer unit tests (e.g., tests/test_red_light_rule.py) "
        "validate synthetic geometry and interval logic—they do not prove field accuracy.",
    )
    _add(
        doc,
        "BodyText",
        "Rule/OCR evaluation artifacts: 30 OCR crops (ocr_review_annotations.csv), 48 helmet candidates "
        "(helmet_reviewed_annotations.csv, all reviewer_label=Unreviewed), one red-light review export (track 130, "
        "11.788 s, unlabeled), zero overtaking predictions in overtaking_predictions_for_review.csv.",
    )

    _add(doc, "Heading1", "Results")
    _add_table_detection(doc, m, n)
    _add_table_rules(doc)
    _add(doc, "Heading2", "Preliminary red-light case study (not a confirmed violation)")
    _add(
        doc,
        "BodyText",
        "On red light.mp4, exploratory processing produced one JSONL candidate for track 130 at 11.788 s "
        "(bbox-center crossing of configured stop line y=540). Repository investigation records: physical painted "
        "stop-bar crossing unconfirmed; ROI frames at the event time show a green straight arrow and countdown, not a "
        "red ball; JSONL signal_state comes from manual interval configuration, not computer vision. The export remains "
        "NEEDS_REVIEW—we do not treat this candidate as a true positive and do not report red-light precision, recall, "
        "or F1.",
    )
    _add(
        doc,
        "BodyText",
        "OCR: EasyOCR returned zero non-empty reads on low-resolution crops (ocr_pipeline_diagnosis.json); CER is "
        "NOT MEASURED until human ground_truth_plate_text is supplied.",
    )

    _add(doc, "Heading1", "Discussion")
    _add(
        doc,
        "BodyText",
        "Detection mAP on 36 test images describes bounding-box quality for annotated classes under CPU inference. "
        "It must not be extrapolated to red-light or overtaking violation accuracy. The track-130 case shows that "
        "center-based crossing can fire in one frame, manual red intervals can disagree with visible green-arrow "
        "phases, and the rule schema does not encode movement-specific aspects—so JSONL candidates can diverge from "
        "visual interpretation even when geometry fires.",
    )

    _add(doc, "Heading1", "Limitations and Threats to Validity")
    _add(
        doc,
        "BodyText",
        "Small test set and high metric variance; train/validation/test stem overlap risk documented in repository "
        "split audits; manual signal intervals without aspect/movement coupling; single-clip red-light investigation; "
        "CPU inference only; YOLOv5+MySQL baseline not implemented in this repository snapshot.",
    )

    _add(doc, "Heading1", "Ethical Considerations")
    _add(
        doc,
        "BodyText",
        "Candidates are not legal findings. Plate and face regions should be redacted in published media; on-chain "
        "storage uses hashes rather than raw plate text where configured.",
    )

    _add(doc, "Heading1", "Future Work")
    _add(
        doc,
        "BodyText",
        "Complete human review packages for OCR, helmet, red-light, and overtaking metrics; aspect-aware signal CSVs "
        "per movement; calibrated overtaking polygons; leak-free dataset splits; optional directional-signal schema "
        "without auto-violation labeling.",
    )

    _add(doc, "Heading1", "Conclusion")
    _add(
        doc,
        "BodyText",
        f"We delivered an evidence-documented prototype with measured object-detection metrics on {n} test images, "
        "passing rule unit tests, documented JSONL candidate generation, and explicit NOT MEASURED violation metrics "
        "until independent human ground truth exists.",
    )

    _add(doc, "Heading5", "Acknowledgment")
    _add(doc, "BodyText", "IT41043 Intelligent Systems coursework prototype; Horizon Campus.")

    _add(doc, "Heading5", "References")
    for i, ref in enumerate(REFERENCES, start=1):
        p = doc.add_paragraph(style="references")
        p.add_run(f"[{i}] ")
        p.add_run(ref)

def validate_docx(path: Path) -> dict:
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    cite_nums = set(int(x) for x in re.findall(r"\[(\d+)\]", text.split("References")[0]))
    ref_nums = set(int(x) for x in re.findall(r"^\[(\d+)\]", text.split("References")[-1], re.M))
    return {
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "char_count": len(text),
        "citations_in_body": sorted(cite_nums),
        "reference_numbers": sorted(ref_nums),
        "citation_audit_ok": cite_nums <= ref_nums and ref_nums <= cite_nums,
        "contains_NOT_MEASURED": "NOT MEASURED" in text,
        "contains_track130_disclaimer": "NEEDS_REVIEW" in text and "11.788" in text,
    }


if __name__ == "__main__":
    build_manuscript()
    print(json.dumps(validate_docx(OUTPUT), indent=2))
