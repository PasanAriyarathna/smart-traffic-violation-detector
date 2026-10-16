"""Generate IEEE-style Markdown and Word draft from verified metrics only."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DETECTOR_DIR = ROOT / "docs/evaluation/detector/detector_val_test_20261009T120646Z"
OUT_MD = ROOT / "docs/research_paper/Smart_Traffic_Violation_IEEE_Paper.md"
OUT_DOCX = ROOT / "docs/research_paper/Smart_Traffic_Violation_IEEE_Paper.docx"


def load_detector() -> dict:
    path = DETECTOR_DIR / "metrics.json"
    if not path.is_file():
        raise SystemExit(f"Missing verified detector metrics: {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def build_markdown(detector: dict) -> str:
    m = detector["metrics"]
    per_class = detector.get("per_class_ap50", [])
    table_rows = "\n".join(
        f"| {row['class']} | {row['ap50']:.3f} |" for row in per_class
    )
    return f"""# Smart Traffic Violation Detection Using YOLOv8, Rule-Based Reasoning, and Human-in-the-Loop Review

**Course / project:** IT41043 Intelligent Systems prototype  
**Evidence date:** {date.today().isoformat()}  
**Git commit (detector run):** `{detector.get('git_commit', 'unknown')}`

> This draft uses **measured** results from `docs/evaluation/`. Rule-level and OCR accuracy are marked NOT MEASURED until human labels exist.

## Abstract

We present a research prototype that detects traffic-related objects with a custom YOLOv8n model, tracks objects in video with ByteTrack, applies geometric and temporal rules for helmet, red-light, and overtaking **candidates**, extracts evidence frames, and supports human review with optional SHA-256 evidence hashing and a Solidity audit logger tested on a local Hardhat network. On a held-out **test split of 36 images** (CPU, conf=0.25, IoU=0.7), the detector achieved precision {m['precision']:.3f}, recall {m['recall']:.3f}, mAP@0.5 {m['map50']:.3f}, and mAP@0.5:0.95 {m['map50_95']:.3f}. These metrics assess **object detection**, not confirmed legal violations. Helmet rule precision and OCR character error rate were **not measured** pending reviewer annotations.

## Index Terms

Traffic violation detection, YOLOv8, ByteTrack, human review, blockchain metadata, dashcam video

## I. Introduction

Automated traffic enforcement from dashcam or fixed cameras requires reliable perception, temporal reasoning, and governance. Our prototype separates **violation candidates** from **verified events** through explicit human review and documents each experiment for reproducibility.

## II. Related Work

Object detection for traffic scenes commonly uses YOLO-family detectors; multi-object tracking (e.g., ByteTrack) stabilizes identities across frames. Legal violation inference requires domain rules beyond bounding-box mAP. We do not claim state-of-the-art detection; we report verifiable metrics on our local dataset.

## III. System Architecture

Pipeline stages: video input → YOLOv8 detection → ByteTrack tracking → rule modules (helmet association, red-light geometry, overtaking lanes) → evidence export (JSONL + images) → FastAPI/SQLite review API → optional chain logging of hashes.

## IV. Dataset and Experimental Setup

- **Local split counts (file audit):** 253 train, 72 validation, 36 test (361 total). See `docs/evaluation/dataset/split_independence_audit.json`.
- **Split independence warning:** Stem-key overlaps exist across splits (potential leakage); generalization claims must be qualified.
- **Detector weights SHA-256:** `{detector['config']['weights_sha256']}`

## V. Results

### A. Object detection (verified)

| Metric | Value |
|--------|------:|
| Precision | {m['precision']:.4f} |
| Recall | {m['recall']:.4f} |
| mAP@0.5 | {m['map50']:.4f} |
| mAP@0.5:0.95 | {m['map50_95']:.4f} |

**Per-class AP@0.5**

| Class | AP@0.5 |
|-------|-------:|
{table_rows}

Source: `{DETECTOR_DIR.as_posix()}/metrics.json`

### B. Helmet rule candidates (partial)

- **48** review candidates verified with evidence images (`helmet_artifact_verification.json`).
- **35** distinct person track IDs.
- Human labels: **0** TP/FP → precision **NOT MEASURED**.

### C. OCR (not measured)

30 plate crops exported; no ground-truth plate text entered.

### D. Red-light and overtaking (not measured on video)

Synthetic unit tests pass; real-video event accuracy **NOT MEASURED**.

### E. Backend API (verified)

Integration tests: import JSONL, list events, verify, chain endpoint returns 503 without RPC config. See `docs/evaluation/backend/api_test_report.md`.

### F. Blockchain (local only)

Hardhat: **2** contract tests passed. Sepolia latency/deployment **NOT MEASURED**.

### G. Baseline (YOLOv5 + MySQL)

**NOT IMPLEMENTED** in this repository.

## VI. Discussion

Detection mAP does not establish violation accuracy. Candidates require human review under a written protocol (`helmet_annotation_guidelines.md`). Dataset split overlap threatens unbiased test evaluation.

## VII. Limitations and Threats to Validity

Small test set, CPU inference, possible train/test leakage, unlabeled rule and OCR evaluation, no public testnet measurements.

## VIII. Ethical and Privacy Considerations

License plates and faces should be redacted in publications; store hashes rather than raw plate text on-chain.

## IX. Conclusion and Future Work

Complete human annotation for helmet and OCR, configure red-light/overtaking on labeled video, fix split leakage, run GPU timing benchmarks, and implement a fair YOLOv5 baseline.

## References

[1] J. Redmon et al., "You Only Look Once: Unified, Real-Time Object Detection," CVPR, 2016.  
[2] G. Jocher et al., Ultralytics YOLO documentation, https://docs.ultralytics.com/  
[3] Y. Zhang et al., "ByteTrack: Multi-Object Tracking by Associating Every Detection Box," ECCV, 2022.
"""


def build_docx(markdown_text: str) -> None:
    try:
        from docx import Document
    except ImportError as exc:
        raise SystemExit("Install python-docx to generate .docx") from exc

    doc = Document()
    for line in markdown_text.splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("|") and "---" not in line:
            doc.add_paragraph(line)
        elif line.strip():
            doc.add_paragraph(line)
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)


def main() -> None:
    detector = load_detector()
    md = build_markdown(detector)
    OUT_MD.parent.mkdir(parents=True, exist_ok=True)
    OUT_MD.write_text(md, encoding="utf-8")
    print(f"Wrote {OUT_MD}")
    try:
        build_docx(md)
        print(f"Wrote {OUT_DOCX}")
    except SystemExit as exc:
        print(exc)


if __name__ == "__main__":
    main()
